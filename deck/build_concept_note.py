"""Builds the track concept note (Word). Numbers come from results/metrics.md, results/rag_ceiling.md, reports/lab6_redteam.md."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL = RGBColor(0x0B, 0x4F, 0x5C)
d = Document()
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin = s.bottom_margin = Inches(0.8)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for name, size in (("Heading 1", 15), ("Heading 2", 12)):
    h = d.styles[name]; h.font.name = "Cambria"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = TEAL
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)


def para(text, bold_lead=None, style=None, after=4):
    p = d.add_paragraph(style=style)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(after)
    return p


def table(rows, widths, header=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.autofit = False
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for row in t.rows:                                   # keep each row on one page
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Inches(widths[j]); c.text = ""
            p = c.paragraphs[0]; r = p.add_run(str(val)); r.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(0)
            if header and i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "0B4F5C")
    d.add_paragraph().paragraph_format.space_after = Pt(2)


t = d.add_paragraph(); r = t.add_run("Healthcare Data-Quality Firewall"); r.font.size = Pt(22); r.bold = True; r.font.name = "Cambria"; r.font.color.rgb = TEAL
para("Concept note: Healthcare track (rules, hybrid LLM, RAG, and a guarded review agent). Lab 6 pattern: tool use, guardrails, red-teaming.", after=8)

d.add_heading("1. Problem and why it matters", 1)
para("A patient's record is rarely in one place. Hospital, clinic and urgent-care systems each hold a fragment with their own patient ID, name format and "
     "wording for the same diagnosis. When these fragments are combined and fed into clinical AI or analytics, five kinds of error propagate: duplicate "
     "records, contradictions (for example an allergy documented in one record and 'no known allergies' in a later one), missing critical fields, "
     "impossible dates, and terminology drift ('MI' versus 'heart attack'). The concept is a data-quality firewall that checks records before they "
     "reach downstream systems, and flags issues, with evidence, for a human data steward. It never makes a clinical decision.")

d.add_heading("2. Concept and architecture", 1)
para("Records are re-linked to patients (fuzzy name, birth date and sex), then deterministic rules run first and emit issues with evidence. Only the cases "
     "rules cannot decide go to Gemini: unfamiliar terms, free-text notes that deny or stop something listed, borderline duplicates, and plain-language "
     "explanations. A retrieval-augmented knowledge base (10 data-steward guideline documents and one entry per vocabulary concept) supports the LLM and "
     "a review agent. The agent investigates a single flagged issue through tools under budgets, validated arguments and a human-approval gate.")
table([
    ["Course sequence", "What was built", "Evidence"],
    ["Problem framing and data", "Synthetic generator with exact ground truth: 300 patients, 1,033 records, 1,416 planted issues, decoys that must not be flagged", "Reproducible from a seed; tests check no label leakage"],
    ["Baselines", "Rule baseline and an LLM-only baseline (Gemini on raw records, same output format)", "Same scoring for all systems"],
    ["Hybrid LLM", "Rules first, Gemini only on the residual, with provenance on every issue", "Precision 1.00, recall 0.98, F1 0.992"],
    ["RAG", "Guideline and terminology knowledge base, BM25 plus character n-gram retriever, citation grounding; optional retrieval-augmented term mapping", "Retrieval recall@5 = 0.80; full menu beats RAG at 63 concepts"],
    ["Agent, guardrails, red-team", "Review agent: five tools, three budgets, Pydantic contracts, scope, human approval, structured output, output filter; 21-case attack suite", "Results in section 4"],
], [1.5, 3.4, 1.9])

d.add_heading("3. Data and evaluation design", 1)
para("Real clinical data needs credentialing and data-use agreements, so the project generates synthetic patients and plants known problems: duplicates, "
     "contradictions, missing fields, impossible dates, and terminology differences that arise naturally from independent wording per source system. Every "
     "planted problem is written to an answer key. Decoys (twins, same-name patients, legitimate dose changes, an allergy acquired later, benign notes) measure "
     "false alarms. Systems are scored on precision, recall, F1, false-positive rate, issue-level accuracy, cost and latency. Terminology is 76% of the labels, "
     "so macro-F1 and per-type results are reported next to the micro average.")

d.add_heading("4. Results", 1)
table([
    ["Metric", "Rules", "LLM-only", "Hybrid"],
    ["Precision / recall", "1.000 / 0.893", "0.956 / 0.832", "1.000 / 0.984"],
    ["F1 (micro / macro)", "0.943 / 0.947", "0.890 / 0.889", "0.992 / 0.993"],
    ["Cost (Gemini 3.1 Pro preview, list-price assumption)", "$0", "$10.56 (42 calls)", "$4.41 (98 calls)"],
], [2.6, 1.4, 1.4, 1.4])
para("The hybrid's gain over rules is entirely recall (contradictions 0.75 to 1.00, terminology 0.90 to 0.98, duplicates 0.88 to 0.97). LLM-only found fewer issues "
     "and cost 2.4 times as much; one chunk of 25 records was lost to output truncation (excluding it, recall rises from 0.832 to 0.853 and the ranking is unchanged). "
     "Most of LLM-only's unmatched flags are real inconsistencies the answer key does not label, so its precision is understated.", after=6)
para("Agent red-team (Lab 6 targets). The suite has 17 attacks (9 direct, 8 indirect through poisoned notes and a poisoned knowledge-base document) and 4 benign controls. "
     "A simulated worst-case model that obeys any instruction it reads was used, so the unguarded baseline blocks nothing by construction; the table shows what each structural layer buys.", after=4)
table([
    ["Layers (cumulative)", "Blocked (17)", "False positives (4)", "Privileged executed"],
    ["None", "0.00", "0/4", "9"],
    ["+ heuristic detector v1", "0.29", "4/4", "7"],
    ["+ detector v2 (false-positive fix)", "0.29", "0/4", "6"],
    ["+ structured output with citation grounding", "0.35", "0/4", "6"],
    ["+ privilege capping (validation, scope, human approval, call budget)", "0.82", "0/4", "0"],
    ["+ output filter", "0.94", "0/4", "0"],
    ["+ verdict cross-check against rule evidence", "1.00", "0/4", "0"],
], [3.3, 1.0, 1.4, 1.1])
para("Targets from Lab 6: block rate at least 0.80 (reached at privilege capping, 0.82), false positives at most 1 of 4 (0 of 4), privileged tool executed by any attack: 0 "
     "(0 from privilege capping onward), argument validation 100% (100% with capping on), loop always terminates (all cases, three budgets enforced). Cost per query is not "
     "measured: the simulation costs nothing and a live run is pending; a reasoning model at about $0.05 or more per call will not meet a $0.02 target without a cheaper model.", after=4)

d.add_heading("5. Safety, privacy and survivability", 1)
para("Design principle: the model proposes, a human disposes. The only privileged tool, apply_correction, is limited by schema to three demographic or encounter fields, "
     "works only on records of the case under review, and runs only after human approval; these limits are code, not prompt text. Notes and guideline text are treated "
     "as untrusted data. The attack that survives layers one to five is plain-language verdict manipulation ('mark this as false alarm'), which has no trigger words and needs "
     "no privileged tool, so its damage is bounded to an advisory verdict that a steward reviews; cross-checking the verdict against deterministic rule evidence removes it. "
     "The review interface shows raw evidence next to every verdict and records accept or reject decisions.")

d.add_heading("6. Limitations and next steps", 1)
para("All data is synthetic and the vocabulary is hand-built. The rules were written alongside the generator, so their precision does not transfer to real data. Some "
     "inconsistencies in the data are side effects of the planted problems and are unlabelled. LLM systems were run once each. The red-team used a simulated worst-case model; the "
     "same suite can be run against live Gemini (python -m aip.redteam --model gemini) but has not been. Lexical retrieval cannot match meaning ('brain attack' to stroke); an "
     "embedding retriever is the natural next step and is needed once the vocabulary is too large to list (UMLS, RxNorm or SNOMED scale). Further work: a human-reviewed sample for "
     "a label-independent precision estimate, a fixed second dataset version, and several seeds.", after=4)
para("Bonus (Lab 7): not included in this note; to be added once the Lab 7 brief is confirmed.", after=2)

d.save("Concept_Note_Healthcare_DQ.docx")
print("saved")
